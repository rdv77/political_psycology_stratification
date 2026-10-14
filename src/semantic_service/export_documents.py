"""Export the shared, source-backed report content to PDF and DOCX.

Run with the bundled Codex Python runtime:
  python export_documents.py ABS_CONTENT_JSON ABS_OUTPUT_DIR
The parent service performs artifact-operation marking before the first run.
"""
from __future__ import annotations

import argparse
import html
import json
import os
import re
import shutil
import subprocess
import sys
import uuid
from pathlib import Path


def _heading(value):
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", str(value), flags=re.UNICODE)).strip()


def _column_weights(headers, rows):
    count = len(headers)
    if not count:
        return []
    weights = []
    for column in range(count):
        lengths = sorted(len(str(row[column])) for row in rows if column < len(row))
        typical = lengths[int(.8 * (len(lengths) - 1))] if lengths else len(str(headers[column]))
        weights.append(min(60., max(10., (max(typical, len(str(headers[column]))) ** .65) * 3)))
    total = sum(weights)
    return [value / total for value in weights]


def _image_path(base, value):
    candidate = (base / str(value)).resolve()
    if not candidate.is_relative_to(base.resolve()):
        raise ValueError(f"Image is outside the report directory: {value}")
    if not candidate.is_file():
        raise FileNotFoundError(f"Report chart is missing: {candidate}")
    return candidate


def _sanitize_docx_title(filename, output):
    """Run the required Google Docs title guard before publishing the DOCX."""
    skill = Path(os.environ.get("SEMANTIC_DOCUMENT_SKILL", "C:/Users/USER/.codex/plugins/cache/"
                                "openai-primary-runtime/documents/26.1013.11914/skills/documents"))
    sanitizer = skill / "scripts/google_docs_title_sanitize.py"
    private = output / ".privatebuild" / "docx"
    private.mkdir(parents=True, exist_ok=True)
    sanitized = private / f"sanitized-{uuid.uuid4().hex}.docx"
    if sanitizer.is_file():
        temporary = private / "tmp"
        temporary.mkdir(parents=True, exist_ok=True)
        environment = {**os.environ, "TMP": str(temporary), "TEMP": str(temporary), "TMPDIR": str(temporary)}
        for command in ([sys.executable, str(sanitizer), str(filename), "--out", str(sanitized)],
                        [sys.executable, str(sanitizer), str(sanitized), "--check"]):
            result = subprocess.run(command, capture_output=True, timeout=120, env=environment)
            if result.returncode:
                message = (result.stderr or result.stdout).decode("utf-8", errors="replace")
                raise RuntimeError("DOCX title sanitation failed: " + message[-1800:])
    else:
        # Portable fallback for a deployment without Codex's bundled skill.
        # Remove exactly title-style/direct title paragraph borders and underlines.
        import zipfile
        from lxml import etree
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        w = "{" + ns["w"] + "}"

        def strip(root):
            for element in list(root.xpath(".//w:pBdr | .//w:u", namespaces=ns)):
                element.getparent().remove(element)

        with zipfile.ZipFile(filename) as source, zipfile.ZipFile(sanitized, "w", zipfile.ZIP_DEFLATED) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename in ("word/styles.xml", "word/document.xml"):
                    root = etree.fromstring(data)
                    if item.filename == "word/styles.xml":
                        for style in root.xpath("//w:style[@w:styleId='Title']", namespaces=ns):
                            strip(style)
                    else:
                        for paragraph in root.xpath("//w:p[w:pPr/w:pStyle[@w:val='Title']]", namespaces=ns):
                            strip(paragraph)
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                target.writestr(item, data)
    shutil.copyfile(sanitized, filename)


def _export_pdf(content, base, output):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Image, KeepTogether, LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    font_dir = Path("C:/Windows/Fonts")
    regular, bold = font_dir / "arial.ttf", font_dir / "arialbd.ttf"
    if not regular.exists() or not bold.exists():
        raise RuntimeError("Cyrillic PDF export requires C:/Windows/Fonts/arial.ttf and arialbd.ttf")
    pdfmetrics.registerFont(TTFont("SemanticArial", str(regular)))
    pdfmetrics.registerFont(TTFont("SemanticArialBold", str(bold)))
    pdfmetrics.registerFontFamily("SemanticArial", normal="SemanticArial", bold="SemanticArialBold")
    styles = getSampleStyleSheet()
    body = ParagraphStyle("ReportBody", parent=styles["BodyText"], fontName="SemanticArial",
                          fontSize=10.5, leading=15, spaceAfter=8, splitLongWords=1,
                          allowWidows=0, allowOrphans=0, textColor=colors.HexColor("#202124"))
    title = ParagraphStyle("ReportTitle", parent=body, fontName="SemanticArialBold", fontSize=24,
                           leading=29, spaceBefore=0, spaceAfter=17, keepWithNext=True, textColor=colors.black)
    section = ParagraphStyle("ReportSection", parent=title, fontSize=16, leading=21,
                             spaceBefore=17, spaceAfter=9)
    sub = ParagraphStyle("ReportSubsection", parent=section, fontSize=12, leading=17,
                         spaceBefore=12, spaceAfter=7)
    caption = ParagraphStyle("ReportCaption", parent=body, fontSize=9, leading=12,
                             spaceAfter=12, textColor=colors.HexColor("#555B63"))
    cell = ParagraphStyle("ReportCell", parent=body, fontSize=9, leading=12, spaceAfter=0)
    header = ParagraphStyle("ReportHeader", parent=cell, fontName="SemanticArialBold", textColor=colors.white)
    source = ParagraphStyle("ReportSource", parent=caption, fontSize=9, leading=12, splitLongWords=1)
    appendix_head = ParagraphStyle("ReportAppendixObject", parent=sub, fontSize=11, leading=14,
                                   spaceBefore=7, spaceAfter=4)
    appendix_body = ParagraphStyle("ReportAppendixBody", parent=body, leading=14, spaceAfter=6)
    page_width, page_height = letter
    margin = 54
    available_width = page_width - 2 * margin
    story = []
    in_source_appendix = False

    def p(text, style=body):
        return Paragraph(html.escape(str(text)).replace("\n", "<br/>"), style)

    for block in content.get("blocks", []):
        kind = block.get("type")
        if kind == "heading":
            level = block.get("level", 1)
            if level == 1:
                in_source_appendix = _heading(block.get("text", "")).casefold() == "объекты и источники"
            style = title if level == 0 else section if level == 1 else appendix_head if in_source_appendix else sub
            flow = p(_heading(block.get("text", "")), style)
            flow._semantic_heading = True
            story.append(flow)
        elif kind == "paragraph":
            story.append(p(block.get("text", ""), appendix_body if in_source_appendix else body))
        elif kind == "table":
            headers, rows = block.get("headers", []), block.get("rows", [])
            if not headers:
                continue
            weights = _column_weights(headers, rows)
            values = [[p(value, header) for value in headers]]
            values.extend([[p(row[index] if index < len(row) else "", cell) for index in range(len(headers))]
                           for row in rows])
            table = LongTable(values, colWidths=[available_width * weight for weight in weights],
                              repeatRows=1, hAlign="LEFT", splitByRow=1, splitInRow=0)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#263C43")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F7F8")]),
                ("GRID", (0, 0), (-1, -1), .5, colors.HexColor("#D9D9D9")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ]))
            story.extend([Spacer(1, 4), table, Spacer(1, 11)])
        elif kind == "image":
            image_path = _image_path(base, block["file"])
            iw, ih = ImageReader(str(image_path)).getSize()
            preceding_headings = []
            while story and getattr(story[-1], "_semantic_heading", False):
                preceding_headings.insert(0, story.pop())
            caption_flow = p(block.get("caption", ""), caption)
            caption_height = caption_flow.wrap(available_width, page_height)[1] + caption.spaceAfter
            heading_height = sum(item.wrap(available_width, page_height)[1] +
                                 item.style.spaceBefore + item.style.spaceAfter for item in preceding_headings)
            available_height = page_height - 48 - 47 - caption_height - heading_height - 22
            ratio = min(available_width / iw, max(120, available_height) / ih)
            image = Image(str(image_path), width=iw * ratio, height=ih * ratio)
            image.hAlign = "CENTER"
            figure = preceding_headings + [Spacer(1, 6), image, caption_flow]
            story.append(KeepTogether(figure))
        elif kind == "link":
            url = str(block.get("url", ""))
            label = str(block.get("text") or url)
            if url.startswith(("https://", "http://")):
                text = f'<link href="{html.escape(url, quote=True)}" color="#176F74">{html.escape(label)}</link>'
                if label != url:
                    text += "<br/>" + html.escape(url)
                story.append(Paragraph(text, source))
            else:
                story.append(p(label, source))
    if not story:
        raise ValueError("The report has no content blocks")

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("SemanticArial", 8.5)
        canvas.setFillColor(colors.HexColor("#60656D"))
        canvas.drawRightString(page_width - margin, 29, f"Страница {document.page}")
        canvas.restoreState()

    document = SimpleDocTemplate(str(output / "report.pdf"), pagesize=letter,
                                 rightMargin=margin, leftMargin=margin, topMargin=48, bottomMargin=47,
                                 title=content.get("study", {}).get("title", "Семантическое пространство"),
                                 author="Semantic Space Research")
    document.build(story, onFirstPage=footer, onLaterPages=footer)


def _export_docx(content, base, output):
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
    from PIL import Image

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    content_width = 7.
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.18
        props = style.element.get_or_add_pPr()
        for border in list(props.findall(qn("w:pBdr"))):
            props.remove(border)
    for name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(16 if name != "Title" else 0)
        style.paragraph_format.keep_with_next = True
    doc.styles["Caption"].font.size = Pt(9)
    doc.styles["Caption"].font.color.rgb = RGBColor.from_string("555B63")

    def add_hyperlink(paragraph, label, url):
        from docx.opc.constants import RELATIONSHIP_TYPE
        relationship = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), relationship)
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "176F74")
        properties.append(color)
        run.append(properties)
        text = OxmlElement("w:t")
        text.text = label
        run.append(text)
        link.append(run)
        paragraph._p.append(link)

    for block in content.get("blocks", []):
        kind = block.get("type")
        if kind == "heading":
            level = block.get("level", 1)
            doc.add_paragraph(_heading(block.get("text", "")), "Title" if level == 0 else "Heading 1" if level == 1 else "Heading 2")
        elif kind == "paragraph":
            doc.add_paragraph(str(block.get("text", "")))
        elif kind == "table":
            headers, rows = block.get("headers", []), block.get("rows", [])
            if not headers:
                continue
            table = doc.add_table(rows=1, cols=len(headers))
            table.autofit = False
            widths = [content_width * weight for weight in _column_weights(headers, rows)]
            for index, width in enumerate(widths):
                table.columns[index].width = Inches(width)
            for index, value in enumerate(headers):
                table.rows[0].cells[index].text = str(value)
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for values in rows:
                cells = table.add_row().cells
                for index in range(len(headers)):
                    cells[index].text = str(values[index]) if index < len(values) else ""
            for row_index, row in enumerate(table.rows):
                for column, cell in enumerate(row.cells):
                    cell.width = Inches(widths[column])
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    props = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement("w:tcBorders")
                    for edge in ("top", "left", "bottom", "right"):
                        border = OxmlElement(f"w:{edge}")
                        border.set(qn("w:val"), "single")
                        border.set(qn("w:sz"), "4")
                        border.set(qn("w:color"), "D9D9D9")
                        borders.append(border)
                    props.append(borders)
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "263C43" if row_index == 0 else "F4F7F8" if row_index % 2 == 0 else "FFFFFF")
                    props.append(shade)
                    margins = OxmlElement("w:tcMar")
                    for edge in ("top", "left", "bottom", "right"):
                        pad = OxmlElement(f"w:{edge}")
                        pad.set(qn("w:w"), "100")
                        pad.set(qn("w:type"), "dxa")
                        margins.append(pad)
                    props.append(margins)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(2)
                        paragraph.paragraph_format.line_spacing = 1.1
                        for run in paragraph.runs:
                            run.font.name = "Arial"
                            run.font.size = Pt(9.5)
                            run.font.bold = row_index == 0
                            run.font.color.rgb = RGBColor(255, 255, 255) if row_index == 0 else RGBColor(0, 0, 0)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        elif kind == "image":
            image_path = _image_path(base, block["file"])
            with Image.open(image_path) as image:
                width, height = image.size
            scale = min(content_width / width, 6.7 / height)
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(image_path), width=Inches(width * scale), height=Inches(height * scale))
            doc.add_paragraph(str(block.get("caption", "")), "Caption")
        elif kind == "link":
            paragraph = doc.add_paragraph()
            label, url = str(block.get("text") or block.get("url", "")), str(block.get("url", ""))
            if url.startswith(("http://", "https://")):
                add_hyperlink(paragraph, label, url)
                if label != url:
                    paragraph.add_run("\n" + url)
            else:
                paragraph.add_run(label)
            for run in paragraph.runs:
                run.font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Страница ").font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = content.get("study", {}).get("title", "Семантическое пространство")
    filename = output / "report.docx"
    doc.save(filename)
    _sanitize_docx_title(filename, output)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("content_json", type=Path)
    parser.add_argument("output_dir", type=Path)
    args = parser.parse_args()
    content_path, output = args.content_json.resolve(), args.output_dir.resolve()
    content = json.loads(content_path.read_text(encoding="utf-8"))
    output.mkdir(parents=True, exist_ok=True)
    _export_pdf(content, content_path.parent, output)
    _export_docx(content, content_path.parent, output)
    print(json.dumps({"pdf": str(output / "report.pdf"), "docx": str(output / "report.docx")}, ensure_ascii=False))


if __name__ == "__main__":
    main()

"""Create an editable Word edition from the political report's shared content.

Requires python-docx and Pillow. No network, scoring changes or paid APIs.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Mm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'research/rwanda_politicians_2021_2026/political_report_q4_2026_10_18'
GITHUB = 'https://github.com/rdv77/political_psycology_stratification/blob/main/research/rwanda_politicians_2021_2026/political_report_q4_2026_10_18/'


def font(style, size, bold=False):
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rf.set(qn('w:'+key), 'Arial')


def hyperlink(paragraph, text, url, bold=False):
    if not url.startswith(('https://', 'http://')):
        url = GITHUB + url
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '167D9A'); props.append(color)
    if bold:
        props.append(OxmlElement('w:b'))
    run.append(props)
    t = OxmlElement('w:t'); t.text = text; run.append(t)
    link.append(run); paragraph._p.append(link)


def rich(paragraph, value):
    tokens = re.split(r'(\*\*.+?\*\*|\[[^\]]+\]\([^)]+\))', value)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            paragraph.add_run(token[2:-2]).bold = True
        else:
            match = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', token)
            if match:
                hyperlink(paragraph, match[1], match[2])
            else:
                paragraph.add_run(token)


def configure(section, landscape=False):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Mm(297 if landscape else 210)
    section.page_height = Mm(210 if landscape else 297)
    section.left_margin = section.right_margin = Mm(15 if landscape else 21)
    section.top_margin = Mm(12 if landscape else 19)
    section.bottom_margin = Mm(16 if landscape else 19)
    section.footer_distance = Mm(8)


def ordered_blocks(blocks):
    result, figures = [], []
    for b in blocks:
        if b['kind'] == 'heading':
            result.extend(figures); figures = []
        if b['kind'] == 'figure':
            figures.append(b)
        else:
            result.append(b)
    return result+figures


def build():
    content = json.loads((OUT/'report_content.json').read_text(encoding='utf-8'))
    doc = Document()
    props = doc.core_properties
    props.title = 'Руанда: власть, развитие и язык политической легитимности'
    props.subject = 'Политологический отчёт по четырём измерениям и пяти группам'
    props.author = 'Political psychology stratification research'
    props.keywords = 'Руанда, политики, публичные позиции'
    props.comments = ''
    configure(doc.sections[0])
    normal = doc.styles['Normal']; font(normal, 11)
    pf = normal.paragraph_format
    pf.line_spacing = 1.12; pf.space_after = Pt(7); pf.widow_control = True
    for name, size in [('Title', 26), ('Subtitle', 13), ('Heading 1', 18), ('Heading 2', 13.5), ('Heading 3', 12)]:
        s = doc.styles[name]; font(s, size, bold=name != 'Subtitle')
        for border in s.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
        s.paragraph_format.space_before = Pt(14 if name.startswith('Heading') else 0)
        s.paragraph_format.space_after = Pt(9)
        s.paragraph_format.keep_with_next = True
    for name, size in [('Caption', 9), ('Quote', 12), ('List Bullet', 11)]:
        font(doc.styles[name], size)
    font(doc.styles['Footer'], 8)
    doc.styles['Quote'].paragraph_format.left_indent = Mm(5)
    # Page numbers remain editable fields and continue across orientation changes.
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    landscape = False
    texts = []

    def paragraph(text, style=None):
        p = doc.add_paragraph(style=style)
        rich(p, text); texts.append(text)
        return p

    def small(text):
        p = paragraph(text)
        p.paragraph_format.space_after = Pt(5)
        for r in p.runs:
            r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string('526672')
        return p

    for block in ordered_blocks(content['blocks']):
        kind, text = block['kind'], block['text']
        if kind != 'figure' and landscape:
            configure(doc.add_section(WD_SECTION.NEW_PAGE)); landscape = False
        if kind == 'title':
            p = paragraph(text, 'Title'); p.paragraph_format.space_before = Mm(15)
        elif kind == 'subtitle':
            paragraph(text, 'Subtitle')
        elif kind == 'heading':
            p = paragraph(text, 'Heading 1')
            if text == 'Основные выводы' or text.startswith('9. '):
                p.paragraph_format.page_break_before = True
        elif kind == 'subheading':
            paragraph(text, 'Heading 2')
        elif kind == 'paragraph':
            paragraph(text)
        elif kind == 'bullet':
            paragraph(text, 'List Bullet')
        elif kind == 'note':
            p = paragraph(text)
            if p.runs:
                # An ordinary paragraph, no box or shading.
                p.runs[0].italic = True
        elif kind == 'quote':
            nested = text.replace('«', '“').replace('»', '”')
            p = paragraph('«'+nested+'»', 'Quote')
            p.paragraph_format.keep_with_next = True
            p = small(block['translation']+'. Оригинал: “'+block['original']+'”')
            p.paragraph_format.keep_with_next = True
            small(block['attribution']+'. ['+block['source_title']+']('+block['source_url']+').')
        elif kind == 'table':
            table = doc.add_table(rows=1, cols=2)
            table.autofit = False
            table.columns[0].width = Mm(49); table.columns[1].width = Mm(119)
            for values in [block['headers']]+block['rows']:
                cells = table.rows[0].cells if values is block['headers'] else table.add_row().cells
                for i, (cell, value) in enumerate(zip(cells, values)):
                    cell.width = Mm(49 if i==0 else 119)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p = cell.paragraphs[0]; rich(p, value)
                    p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.08
                    for r in p.runs: r.font.size = Pt(10)
                    tcpr = cell._tc.get_or_add_tcPr()
                    margins = OxmlElement('w:tcMar')
                    for edge, size in [('top',100), ('bottom',100), ('left',110), ('right',110)]:
                        el = OxmlElement('w:'+edge); el.set(qn('w:w'), str(size)); el.set(qn('w:type'), 'dxa'); margins.append(el)
                    tcpr.append(margins)
                no_split = OxmlElement('w:cantSplit'); cells[0]._tc.getparent().get_or_add_trPr().append(no_split)
            header = table.rows[0]
            repeat = OxmlElement('w:tblHeader'); header._tr.get_or_add_trPr().append(repeat)
            for cell in header.cells:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'),'EAF1F4'); cell._tc.get_or_add_tcPr().append(shade)
                for r in cell.paragraphs[0].runs: r.bold = True
            borders = OxmlElement('w:tblBorders')
            for edge in ('top','left','bottom','right','insideH','insideV'):
                el = OxmlElement('w:'+edge); el.set(qn('w:val'),'single');el.set(qn('w:sz'),'5');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            table._tbl.tblPr.append(borders)
            paragraph('')
        elif kind == 'figure':
            configure(doc.add_section(WD_SECTION.NEW_PAGE), landscape=True); landscape = True
            paragraph(text, 'Heading 2')
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(3); p.paragraph_format.keep_with_next = True
            path = OUT/block['path']; w,h = Image.open(path).size
            # Up to 267 mm wide, with room for title and caption on A4 landscape.
            width = min(267, 153*w/h)
            shape = p.add_run().add_picture(str(path), width=Mm(width))
            shape._inline.docPr.set('descr', text+'. '+block['caption'])
            paragraph(block['caption'], 'Caption')
        elif kind == 'biography':
            paragraph(text, 'Heading 3')
            p = paragraph(block['body']);p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True
            p = small('Период/дата: '+block['period']+'. '+block['group']);p.paragraph_format.keep_with_next = True
            p.paragraph_format.keep_together = True
            refs = '; '.join('['+s['title']+']('+s['url']+')' for s in block['sources'])
            small('Источники: '+refs+'.')

    path = OUT/'REPORT.docx'
    doc.save(path)
    # Content audit includes text in hyperlinks and cells, not just doc.paragraphs.
    from lxml import etree
    with ZipFile(path) as archive:
        xml = etree.fromstring(archive.read('word/document.xml'))
        ns = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        all_text = '\n'.join(xml.xpath('//w:t/text()', namespaces=ns))
        paragraph_text = {' '.join(''.join(p.xpath('.//w:t/text()',namespaces=ns)).split())
                          for p in xml.xpath('//w:body//w:p',namespaces=ns)}
        images = [n for n in archive.namelist() if n.startswith('word/media/')]
    profiles = json.loads((OUT/'biographies.json').read_text(encoding='utf-8'))['profiles']
    assert all(p['name_original'] in all_text for p in profiles)
    assert len(images) == 4
    quote_count = sum(b['kind']=='quote' for b in content['blocks'])
    assert quote_count == 7
    assert all(b['body'] in all_text for b in content['blocks'] if b['kind']=='biography')
    checked = 0
    for b in content['blocks']:
        if b['kind'] in ('paragraph','bullet','note'):
            text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1',b['text']).replace('**','')
            assert ' '.join(text.split()) in paragraph_text
            checked += 1
    assert sum(s.orientation == WD_ORIENT.LANDSCAPE for s in doc.sections)==4
    receipt = {'filename':'REPORT.docx', 'participants':42, 'embedded_diagrams':len(images),
               'quotes':quote_count, 'landscape_sections':4, 'tables':len(doc.tables),
               'all_biographies_preserved':True, 'narrative_paragraphs_preserved':checked, 'editable_text':True,
               'source_content_sha256':hashlib.sha256((OUT/'report_content.json').read_bytes()).hexdigest(),
               'docx_sha256':hashlib.sha256(path.read_bytes()).hexdigest(), 'new_paid_api_requests':0}
    (OUT/'docx_validation.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(receipt,ensure_ascii=False))


if __name__ == '__main__':
    build()
